"""Generadores de PDF y DOCX de prueba, en memoria."""

import io
import zipfile

from docx import Document
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter


def make_pdf(*pages: str) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for text in pages or ("",):
        pdf.add_page()
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def make_encrypted_pdf(text: str) -> bytes:
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(make_pdf(text))))
    writer.encrypt("secreto")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()
