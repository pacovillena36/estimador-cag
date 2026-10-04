"""Adjuntos de una conversación (PDF y DOCX): validación y extracción de texto.

Camino B (extracción local con pypdf y python-docx): independiente del
proveedor LLM, con control total sobre tamaño y contenido, y prepara el
chunking de RAG de módulos posteriores.

Todo se procesa en memoria (BytesIO), sin escribir a disco. Las
comprobaciones, en orden:
1. Número de archivos y tamaño (por archivo y total), leyendo en bloques y
   cortando al superar el límite: no se confía en Content-Length.
2. Tipo real: extensión .pdf/.docx Y firma de bytes (%PDF- / PK\\x03\\x04).
   No se confía en el content_type que declara el cliente.
3. DOCX (un ZIP): antes de abrirlo, número de entradas y suma de tamaños
   descomprimidos (protección contra zip bombs).
4. PDF: cifrados o corruptos se rechazan; solo se procesan las primeras
   páginas.
5. Texto extraído: límite de caracteres por adjunto y total, con marca de
   truncado, y neutralización de las etiquetas delimitadoras del prompt.

Las funciones de extracción son bloqueantes: el servicio las ejecuta en un
threadpool para no bloquear el event loop.
"""

import io
import unicodedata
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import PureWindowsPath
from typing import Literal

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

from app.prompts.loader import neutralize_session_delimiters as neutralize_delimiters

TRUNCATION_MARK = "\n[... contenido truncado ...]"
MAX_FILENAME_CHARS = 100
_READ_CHUNK = 64 * 1024
_SIGNATURES = {"pdf": b"%PDF-", "docx": b"PK\x03\x04"}

AttachmentKind = Literal["pdf", "docx"]


@dataclass(frozen=True)
class AttachmentLimits:
    max_files: int
    max_bytes: int
    max_total_bytes: int
    max_pdf_pages: int
    max_docx_uncompressed_bytes: int
    max_docx_entries: int
    max_chars: int
    max_total_chars: int

    @classmethod
    def from_settings(cls, settings) -> "AttachmentLimits":
        return cls(
            max_files=settings.max_attachments,
            max_bytes=settings.max_attachment_bytes,
            max_total_bytes=settings.max_total_attachment_bytes,
            max_pdf_pages=settings.max_pdf_pages,
            max_docx_uncompressed_bytes=settings.max_docx_uncompressed_bytes,
            max_docx_entries=settings.max_docx_entries,
            max_chars=settings.max_attachment_chars,
            max_total_chars=settings.max_total_attachment_chars,
        )


@dataclass(frozen=True)
class RawAttachment:
    filename: str  # ya saneado
    kind: AttachmentKind
    data: bytes


@dataclass(frozen=True)
class ExtractedAttachment:
    filename: str
    text: str
    truncated: bool


# ----------------------------------------------------------------- errores


class AttachmentError(Exception):
    """Error de validación de un adjunto. `message` es genérico y apto para
    el cliente; `reason` es solo para logs."""

    status_code = 422
    code = "invalid_attachment"
    message = "No se ha podido procesar uno de los adjuntos."

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


class AttachmentTooLargeError(AttachmentError):
    status_code = 413
    code = "attachment_too_large"
    message = "Los adjuntos superan el tamaño o el número máximo permitido."


class UnsupportedAttachmentError(AttachmentError):
    status_code = 415
    code = "unsupported_attachment"
    message = "Solo se admiten adjuntos PDF (.pdf) y Word (.docx)."


class UnreadableAttachmentError(AttachmentError):
    status_code = 422
    code = "unreadable_attachment"
    message = "Uno de los adjuntos está cifrado o dañado y no se puede leer."


# --------------------------------------------------------- funciones puras


def sanitize_filename(name: str | None) -> str:
    """Solo el nombre base, sin caracteres de control ni saltos de línea y
    con longitud acotada (se usa en el prompt y en logs)."""
    base = PureWindowsPath(name or "").name  # separa por "/" y por "\\"
    cleaned = "".join(ch for ch in base if not unicodedata.category(ch).startswith("C"))
    cleaned = cleaned.strip()[:MAX_FILENAME_CHARS]
    return cleaned or "adjunto"


def detect_kind(filename: str, head: bytes) -> AttachmentKind:
    """Extensión y firma de bytes deben coincidir."""
    extension = PureWindowsPath(filename).suffix.lower()
    kind: AttachmentKind | None = {".pdf": "pdf", ".docx": "docx"}.get(extension)  # type: ignore[assignment]
    if kind is None:
        raise UnsupportedAttachmentError(f"extension:{extension or 'none'}")
    if not head.startswith(_SIGNATURES[kind]):
        raise UnsupportedAttachmentError(f"signature_mismatch:{kind}")
    return kind


def _truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars] + TRUNCATION_MARK, True


def extract_pdf_text(data: bytes, *, max_pages: int) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadableAttachmentError("pdf_encrypted")
        pages = reader.pages[:max_pages]
        return "\n".join((page.extract_text() or "").strip() for page in pages).strip()
    except AttachmentError:
        raise
    except Exception as exc:  # pypdf lanza tipos variados ante PDFs malformados
        raise UnreadableAttachmentError(f"pdf_unreadable:{type(exc).__name__}") from exc


def check_docx_archive(data: bytes, *, max_uncompressed_bytes: int, max_entries: int) -> None:
    """Inspecciona el ZIP sin descomprimirlo (protección contra zip bombs)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise UnreadableAttachmentError("docx_bad_zip") from exc
    if len(entries) > max_entries:
        raise AttachmentTooLargeError(f"docx_entries:{len(entries)}")
    if sum(entry.file_size for entry in entries) > max_uncompressed_bytes:
        raise AttachmentTooLargeError("docx_uncompressed_size")
    if not any(entry.filename == "word/document.xml" for entry in entries):
        raise UnsupportedAttachmentError("docx_not_word")


def extract_docx_text(data: bytes, *, max_uncompressed_bytes: int, max_entries: int) -> str:
    check_docx_archive(data, max_uncompressed_bytes=max_uncompressed_bytes, max_entries=max_entries)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # XML o relaciones dañadas
        raise UnreadableAttachmentError(f"docx_unreadable:{type(exc).__name__}") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_attachments(
    raws: Sequence[RawAttachment], limits: AttachmentLimits
) -> list[ExtractedAttachment]:
    """Bloqueante (ejecutar en threadpool). Aplica los límites de caracteres
    por adjunto y total, y neutraliza las etiquetas delimitadoras."""
    extracted: list[ExtractedAttachment] = []
    remaining = limits.max_total_chars
    for raw in raws:
        if raw.kind == "pdf":
            text = extract_pdf_text(raw.data, max_pages=limits.max_pdf_pages)
        else:
            text = extract_docx_text(
                raw.data,
                max_uncompressed_bytes=limits.max_docx_uncompressed_bytes,
                max_entries=limits.max_docx_entries,
            )
        text, truncated = _truncate(neutralize_delimiters(text), min(limits.max_chars, max(remaining, 0)))
        remaining -= len(text)
        extracted.append(ExtractedAttachment(raw.filename, text, truncated))
    return extracted


# -------------------------------------------------------- lectura (async)


async def read_attachments(
    uploads: Sequence[UploadFile], limits: AttachmentLimits
) -> list[RawAttachment]:
    """Lee los archivos en bloques, cortando en cuanto se supera un límite,
    y valida el tipo real. Sin escritura a disco."""
    uploads = [u for u in uploads if u.filename or u.size]  # ignora campos vacíos
    if len(uploads) > limits.max_files:
        raise AttachmentTooLargeError(f"too_many_files:{len(uploads)}")

    raws: list[RawAttachment] = []
    total = 0
    for upload in uploads:
        filename = sanitize_filename(upload.filename)
        buffer = bytearray()
        while chunk := await upload.read(_READ_CHUNK):
            buffer.extend(chunk)
            if len(buffer) > limits.max_bytes:
                raise AttachmentTooLargeError("file_too_large")
            if total + len(buffer) > limits.max_total_bytes:
                raise AttachmentTooLargeError("total_too_large")
        total += len(buffer)
        kind = detect_kind(filename, bytes(buffer[:8]))
        raws.append(RawAttachment(filename, kind, bytes(buffer)))
    return raws


def compose_user_message(transcript: str, attachments: Sequence[ExtractedAttachment]) -> str:
    """Mensaje de usuario del turno: la transcripción y cada adjunto en su
    bloque delimitado (el system prompt indica que son datos, no
    instrucciones)."""
    parts = [f"<transcript>\n{neutralize_delimiters(transcript)}\n</transcript>"]
    for attachment in attachments:
        parts.append(
            f"--- attachment: {attachment.filename} ---\n"
            f"<attachment_content>\n{attachment.text}\n</attachment_content>"
        )
    return "\n\n".join(parts)
